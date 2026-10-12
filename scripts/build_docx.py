#!/usr/bin/env python
"""Build submission-style Word files from the markdown manuscript.

    unset PYTHONPATH; source env/venv/bin/activate
    python scripts/build_docx.py            # writes manuscript/manuscript.docx and manuscript/supplementary.docx

Inputs : manuscript/manuscript.md, manuscript/figure_plan.md, figures/final/Fig*.png
Outputs: manuscript/manuscript.docx   (main text, tables, references, main figure legends + Figs 1-4 at the end)
         manuscript/supplementary.docx (Figures S1-S12 that exist, with legends from figure_plan.md)

Formatting: Times New Roman 12 pt, double-spaced body, 1-inch margins, continuous line numbers
(sectPr/lnNumType), page numbers centred in the footer, real Word tables, bold/italic/code/superscript
from the markdown. Figures at 6.5-inch width. Idempotent: outputs are rebuilt from scratch on every run
with fixed document metadata, so reruns after edits give the same result for the same inputs.

Markdown subset handled (the subset used by manuscript.md): YAML front matter, #/##/### headings,
paragraphs, '---' rules (dropped), pipe tables, '- ' bullets, '1. ' numbered items, **bold**, *italic*,
`code`, ^superscript^ (pandoc style), x^{...} superscripts, and backslash-escaped '*'.
"""
from __future__ import annotations

import datetime as _dt
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
MS_MD = ROOT / "manuscript" / "manuscript.md"
PLAN_MD = ROOT / "manuscript" / "figure_plan.md"
FIG_DIR = ROOT / "figures" / "final"
OUT_MAIN = ROOT / "manuscript" / "manuscript.docx"
OUT_SUPP = ROOT / "manuscript" / "supplementary.docx"

FONT = "Times New Roman"
MONO = "Courier New"
BODY_PT = 12
TABLE_PT = 9
FIG_WIDTH_IN = 6.5
FIG_MAX_HEIGHT_IN = 8.6  # keep very tall figures on one page (width shrinks only if needed)
ESC_STAR = "\ue000"  # placeholder for a backslash-escaped '*'

# --------------------------------------------------------------------------- document setup


def _set_rfonts(rpr, name: str) -> None:
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def _style_font(style, size_pt=None, bold=None, italic=None, name=FONT):
    style.font.name = name
    _set_rfonts(style.element.get_or_add_rPr(), name)
    if size_pt is not None:
        style.font.size = Pt(size_pt)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    style.font.color.rgb = RGBColor(0, 0, 0)


def _add_page_number_footer(section) -> None:
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for kind, text in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text == " PAGE ":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)
    run.font.name = FONT
    run.font.size = Pt(10)


def _add_line_numbers(section) -> None:
    sectPr = section._sectPr
    for old in sectPr.findall(qn("w:lnNumType")):
        sectPr.remove(old)
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:distance"), "360")  # 0.25 inch from text
    ln.set(qn("w:restart"), "continuous")
    # schema order: lnNumType must precede pgNumType/cols/docGrid etc.; insert before the first of these
    successors = ("w:pgNumType", "w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg",
                  "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid", "w:printerSettings")
    for tag in successors:
        el = sectPr.find(qn(tag))
        if el is not None:
            el.addprevious(ln)
            return
    sectPr.append(ln)


def new_document(title: str, date: str) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(1))
    _add_line_numbers(sec)
    _add_page_number_footer(sec)

    st = doc.styles
    _style_font(st["Normal"], BODY_PT)
    pf = st["Normal"].paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before, pf.space_after = Pt(0), Pt(0)
    _style_font(st["Title"], 16, bold=True)
    tpf = st["Title"].paragraph_format
    tpf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tpf.space_after = Pt(12)
    tpf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    # remove the default Title bottom border
    ppr = st["Title"].element.get_or_add_pPr()
    for b in ppr.findall(qn("w:pBdr")):
        ppr.remove(b)
    for lvl, size in ((1, 14), (2, 12), (3, 12)):
        h = st[f"Heading {lvl}"]
        _style_font(h, size, bold=True, italic=(lvl == 3))
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(0)
        h.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        h.paragraph_format.keep_with_next = True
    for name in ("Table Grid",):
        _style_font(st[name], TABLE_PT)

    cp = doc.core_properties
    cp.title = title
    cp.author = "Tharun Nambi"
    cp.comments = "Generated by scripts/build_docx.py from manuscript markdown"
    try:
        stamp = _dt.datetime.strptime(date, "%Y-%m-%d")
    except ValueError:
        stamp = _dt.datetime(2026, 1, 1)
    cp.created = cp.modified = stamp
    cp.last_modified_by = "scripts/build_docx.py"
    cp.revision = 1
    return doc


# --------------------------------------------------------------------------- inline markdown

# order matters: code first (its content is literal), then bold, italic, superscripts
_TOKEN = re.compile(
    r"(?P<code>`[^`]+`)"
    r"|(?P<bold>\*\*(?=\S)(?:.+?)(?<=\S)\*\*)"
    r"|(?P<ital>(?<![\w*])\*(?=[^\s*])(?:[^*]+?)(?<=[^\s*])\*(?![\w*]))"
    r"|(?P<sup>\^(?=\S)[^^\s]+?\^)"
    r"|(?P<supb>\^\{[^}]+\})"
)


def _prep(text: str) -> str:
    return text.replace("\\*", ESC_STAR).replace("\\_", "_").replace("\\#", "#")


def _emit(par, text, bold=False, italic=False, sup=False, mono=False, size=None):
    if not text:
        return
    run = par.add_run(text.replace(ESC_STAR, "*"))
    run.bold = bold or None
    run.italic = italic or None
    if sup:
        run.font.superscript = True
    if mono:
        run.font.name = MONO
        _set_rfonts(run._r.get_or_add_rPr(), MONO)
        run.font.size = Pt((size or BODY_PT) - 1)
    elif size:
        run.font.size = Pt(size)


def add_inline(par, text: str, bold=False, italic=False, size=None) -> None:
    """Add markdown-formatted inline text to a paragraph as runs."""
    text = _prep(text) if ESC_STAR not in text else text
    pos = 0
    for m in _TOKEN.finditer(text):
        _emit(par, text[pos:m.start()], bold, italic, size=size)
        tok = m.group(0)
        if m.group("code"):
            _emit(par, tok[1:-1], bold, italic, mono=True, size=size)
        elif m.group("bold"):
            add_inline(par, tok[2:-2], True, italic, size)
        elif m.group("ital"):
            add_inline(par, tok[1:-1], bold, True, size)
        elif m.group("sup"):
            _emit(par, tok[1:-1], bold, italic, sup=True, size=size)
        elif m.group("supb"):
            _emit(par, tok[2:-1], bold, italic, sup=True, size=size)
        pos = m.end()
    _emit(par, text[pos:], bold, italic, size=size)


# --------------------------------------------------------------------------- block helpers


def add_para(doc, text, style=None, align=None, keep_next=False, indent_hanging=None, prefix=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if keep_next:
        p.paragraph_format.keep_with_next = True
    if indent_hanging is not None:
        p.paragraph_format.left_indent = Inches(indent_hanging)
        p.paragraph_format.first_line_indent = Inches(-indent_hanging)
    if prefix:
        p.add_run(prefix)
    add_inline(p, text)
    return p


def _split_row(line: str) -> list[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def _set_cell_margins(table, twips=60):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(twips if side in ("left", "right") else 20))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def add_table(doc, rows: list[list[str]], warnings: list[str], label: str):
    header, body = rows[0], rows[1:]
    ncol = len(header)
    for r in body:
        if len(r) != ncol:
            warnings.append(f"{label}: row has {len(r)} cells, header {ncol}: {r[:2]}")
    table = doc.add_table(rows=1 + len(body), cols=ncol)
    table.style = doc.styles["Table Grid"]
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    _set_cell_margins(table)
    for i, r in enumerate([header] + body):
        cells = table.rows[i].cells
        for j in range(ncol):
            txt = r[j] if j < len(r) else ""
            if j == ncol - 1 and len(r) > ncol:  # cell text containing '|' (e.g. |rho|): rejoin
                txt = " | ".join(r[j:])
            cell = cells[j]
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, txt, bold=(i == 0), size=TABLE_PT)
    _repeat_header(table.rows[0])
    # small single-spaced gap after the table
    gap = doc.add_paragraph()
    gap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    return table


def add_figure(doc, png: Path):
    from PIL import Image  # pillow is in the project venv (matplotlib dependency)

    w, h = Image.open(png).size
    width = FIG_WIDTH_IN
    if width * h / w > FIG_MAX_HEIGHT_IN:
        width = FIG_MAX_HEIGHT_IN * w / h
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.add_run().add_picture(str(png), width=Inches(width))
    return p, width


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# --------------------------------------------------------------------------- markdown parsing


def parse_front_matter(lines):
    meta, i = {}, 0
    if lines and lines[0].strip() == "---":
        for i in range(1, len(lines)):
            if lines[i].strip() == "---":
                break
            m = re.match(r'(\w+):\s*"?(.*?)"?\s*$', lines[i])
            if m:
                meta[m.group(1)] = m.group(2)
        return meta, lines[i + 1:]
    return meta, lines


def blocks(lines):
    """Yield (kind, payload) blocks from markdown lines."""
    i, n = 0, len(lines)
    while i < n:
        line = lines[i].rstrip("\n")
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith("```"):
            j = i + 1
            while j < n and not lines[j].strip().startswith("```"):
                j += 1
            yield "code", [l.rstrip("\n") for l in lines[i + 1:j]]
            i = j + 1
            continue
        if re.fullmatch(r"-{3,}|\*{3,}", s):
            yield "hr", None
            i += 1
            continue
        m = re.match(r"(#{1,6})\s+(.*)", s)
        if m:
            yield f"h{len(m.group(1))}", m.group(2).strip()
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                r = _split_row(lines[i])
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in r):
                    rows.append(r)
                i += 1
            yield "table", rows
            continue
        m = re.match(r"[-*+]\s+(.*)", s)
        if m:
            yield "bullet", m.group(1)
            i += 1
            continue
        m = re.match(r"(\d+)\.\s+(.*)", s)
        if m:
            yield "num", (m.group(1), m.group(2))
            i += 1
            continue
        # paragraph: consecutive non-blank lines that are not other block starts
        buf = [s]
        i += 1
        while i < n:
            t = lines[i].strip()
            if not t or t.startswith(("#", "|", "```")) or re.match(r"([-*+]|\d+\.)\s+", t) \
                    or re.fullmatch(r"-{3,}", t):
                break
            buf.append(t)
            i += 1
        yield "para", buf


# --------------------------------------------------------------------------- main manuscript


def build_main(warnings):
    raw = MS_MD.read_text(encoding="utf-8").splitlines()
    meta, lines = parse_front_matter(raw)
    title = meta.get("title", "Manuscript")
    doc = new_document(title, meta.get("draft_date", "2026-01-01"))

    # split off the "Figure legends" section so legends + figures go at the end, after the references
    legend_start = next((k for k, l in enumerate(lines) if re.match(r"##\s+Figure legends", l)), None)
    legend_lines = []
    if legend_start is not None:
        legend_end = next((k for k in range(legend_start + 1, len(lines)) if re.match(r"##\s", lines[k])),
                          len(lines))
        legend_lines = lines[legend_start:legend_end]
        lines = lines[:legend_start] + lines[legend_end:]

    render_blocks(doc, lines, warnings, "manuscript")

    # Figure legends + figures at the end
    legends = []
    for kind, payload in blocks(legend_lines):
        if kind == "para":
            text = " ".join(payload)
            m = re.match(r"\*\*Figure (\d+)\.", text)
            if m:
                legends.append((int(m.group(1)), text))
    n_fig = 0
    if legends:
        page_break(doc)
        doc.add_heading("Figure legends", level=1)
        for num, text in legends:
            add_para(doc, text)
        for num, text in legends:
            png = FIG_DIR / f"Fig{num}.png"
            page_break(doc)
            add_para(doc, text, keep_next=True)
            if png.exists():
                add_figure(doc, png)
                n_fig += 1
            else:
                warnings.append(f"missing {png.relative_to(ROOT)}")
                add_para(doc, f"[Figure {num} file missing: {png.relative_to(ROOT)}]")
    OUT_MAIN.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_MAIN)
    return n_fig


def render_blocks(doc, lines, warnings, label):
    table_no = 0
    in_refs = False
    for kind, payload in blocks(lines):
        if kind == "hr":
            continue
        if kind == "h1":
            add_para(doc, payload, style="Title", align=WD_ALIGN_PARAGRAPH.CENTER)
        elif kind == "h2":
            in_refs = payload.lower().startswith("references")
            add_para(doc, payload, style="Heading 1")
        elif kind == "h3":
            add_para(doc, payload, style="Heading 2")
        elif kind in ("h4", "h5", "h6"):
            add_para(doc, payload, style="Heading 3")
        elif kind == "table":
            table_no += 1
            add_table(doc, payload, warnings, f"{label} table {table_no}")
        elif kind == "bullet":
            add_para(doc, payload, indent_hanging=0.25, prefix="•\t")
        elif kind == "num":
            num, text = payload
            add_para(doc, text, indent_hanging=0.35 if in_refs else 0.3, prefix=f"{num}.\t")
        elif kind == "code":
            for l in payload:
                p = doc.add_paragraph()
                _emit(p, l, mono=True)
        elif kind == "para":
            # pandoc-style hard line breaks: lines in one paragraph that start with ^n^ (affiliations)
            if len(payload) > 1 and all(re.match(r"\^", t) for t in payload):
                for t in payload:
                    add_para(doc, t)
                continue
            text = " ".join(payload)
            is_caption = bool(re.match(r"\*\*Table \d+\.", text))
            add_para(doc, text, keep_next=is_caption)


# --------------------------------------------------------------------------- supplementary


def parse_supp_legends():
    legends = {}
    for line in PLAN_MD.read_text(encoding="utf-8").splitlines():
        s = line.strip()
        m = re.match(r"\|\s*S(\d+)\s*\|", s)
        if not m:
            continue
        inner = s[1:-1] if s.endswith("|") else s[1:]
        # columns: # | Data | Status | Legend; the legend may itself contain '|' (e.g. |rho|)
        parts = inner.split("|", 3)
        if len(parts) == 4:
            legends[int(m.group(1))] = (parts[2].strip(), parts[3].strip())
    return legends


def build_supp(warnings):
    raw = MS_MD.read_text(encoding="utf-8").splitlines()
    meta, _ = parse_front_matter(raw)
    title = meta.get("title", "Manuscript")
    doc = new_document("Supplementary Figures: " + title, meta.get("draft_date", "2026-01-01"))
    add_para(doc, "Supplementary Figures", style="Title", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, title, align=WD_ALIGN_PARAGRAPH.CENTER)
    legends = parse_supp_legends()
    included, missing = [], []
    for n in range(1, 13):
        status, legend = legends.get(n, ("", ""))
        png = FIG_DIR / f"FigS{n}.png"
        if not png.exists():
            missing.append(n)
            continue
        page_break(doc)
        add_para(doc, f"Figure S{n}", style="Heading 1")
        add_para(doc, f"**Figure S{n}.** {legend}" if legend else f"**Figure S{n}.**", keep_next=True)
        if not legend:
            warnings.append(f"no legend for S{n} in figure_plan.md")
        add_figure(doc, png)
        included.append(n)
    if missing:
        page_break(doc)
        add_para(doc, "Figures not included", style="Heading 1")
        for n in missing:
            status, legend = legends.get(n, ("", ""))
            add_para(doc, f"**Figure S{n}** is not yet drawn ({status or 'no file'}). Planned legend: {legend}")
    doc.save(OUT_SUPP)
    return included, missing


def main():
    warnings: list[str] = []
    n_main = build_main(warnings)
    inc, miss = build_supp(warnings)
    print(f"wrote {OUT_MAIN.relative_to(ROOT)} ({n_main} main figures)")
    print(f"wrote {OUT_SUPP.relative_to(ROOT)} (Figures S{', S'.join(map(str, inc))}; missing: {miss or 'none'})")
    for w in warnings:
        print("WARNING:", w, file=sys.stderr)


if __name__ == "__main__":
    main()
